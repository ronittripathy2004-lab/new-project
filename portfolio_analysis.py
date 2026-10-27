#!/usr/bin/env python3
"""
Factsheet Automation - Top 10 Holdings & Market-Cap Split
=========================================================

Reads
  * every Excel file inside the PORTFOLIO folder (one monthly portfolio per fund;
    file names are random, so the fund is identified from the TITLE inside the
    file using the FUND_RULES table below)
  * reference.xlsx (AMFI / SEBI market-cap categorisation, matched on ISIN)

Writes (both in the Factsheet Automation folder)
  1. Top10_Holdings.xlsx   - one sheet per fund (ALL funds): top 10 holdings
                             with name and % to net assets
  2. Market_Cap_Split.xlsx - one sheet per EQUITY fund: Large / Mid / Small cap
                             split; Cash & Others = 100 - (Large + Mid + Small)

Rules
  * A holding's cap category comes from matching its ISIN in the reference file.
  * Equity holdings that are NOT found in the reference file (e.g. new listings)
    are listed at the bottom of that fund's split sheet and, per the
    "100 - rest = cash" rule, fall into Cash & Others.
  * Top 10 uses the "% to Net Assets" column. Sub Total / Total / Grand Total /
    Net Receivables rows are ignored, and so are negative rows (short
    derivative / hedge positions).
"""

import os
import re
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# =============================================================================
# 1. CONFIG  -  CHANGE PATHS / TABLES HERE
# =============================================================================

BASE_DIR = "/Users/ronittripathy/Desktop/Factsheet Automation"   # folder that holds reference.xlsx and portfolio/

REFERENCE_FILE = os.path.join(BASE_DIR, "reference.xlsx")
PORTFOLIO_DIR  = os.path.join(BASE_DIR, "portfolio")             # one Excel file per fund
TOP10_OUTPUT   = os.path.join(BASE_DIR, "Top10_Holdings.xlsx")
SPLIT_OUTPUT   = os.path.join(BASE_DIR, "Market_Cap_Split.xlsx")

TOP_N = 10                      # number of holdings to show
TOTAL_ASSETS_PCT = 100          # Cash & Others = this - (Large + Mid + Small)
STRIP_FROM_SHEET_NAMES = ["JioBlackRock "]   # removed from fund titles to make shorter sheet names

# -----------------------------------------------------------------------------
# FUND CLASSIFICATION TABLE   (ORDER MATTERS - first match wins)
#   (display name, fund type, [keywords])
#   A fund matches a row if ANY keyword is found in the fund's title inside the
#   portfolio file (spaces, dashes and capital letters are ignored).
#   If the title matches nothing, the sheet name and then the file name are tried.
#   Put more specific rows ABOVE general ones (e.g. "Ultra Short" above "Short Term").
#
#   Fund types:  "EQUITY"  -> gets the Large/Mid/Small/Cash split
#                "HYBRID"  -> top 10 only (see SPLIT_FUND_TYPES to change)
#                "DEBT"    -> top 10 only
# -----------------------------------------------------------------------------
FUND_RULES = [
    ("Nifty 50 ETF",                    "EQUITY", ["nifty50etf", "nifty50exchangetraded"]),
    ("Balanced Advantage Fund",         "HYBRID", ["balancedadvantage"]),
    ("Arbitrage Fund",                  "HYBRID", ["arbitrage"]),
    ("Large Cap Fund",                  "EQUITY", ["largecap"]),
    ("Flexi Cap Fund",                  "EQUITY", ["flexicap", "jbflexi"]),
    ("Sector Rotation Fund",            "EQUITY", ["sectorrotation"]),
    ("Nifty Smallcap 250 Index Fund",   "EQUITY", ["smallcap"]),
    ("Nifty Midcap 150 Index Fund",     "EQUITY", ["midcap"]),
    ("Nifty Next 50 Index Fund",        "EQUITY", ["next50"]),
    ("Nifty 50 Index Fund",             "EQUITY", ["nifty50index"]),
    ("Nifty 8-13 yr G-Sec Index Fund",  "DEBT",   ["813yr", "gsec"]),
    ("Ultra Short to Short Term Fund",  "DEBT",   ["ultrashort"]),
    ("Short Term Fund",                 "DEBT",   ["shortterm"]),
    ("Money Market Fund",               "DEBT",   ["moneymarket"]),
    ("Overnight Fund",                  "DEBT",   ["overnight"]),
    ("Liquid Fund",                     "DEBT",   ["liquid"]),
]

# Which fund types get the Large/Mid/Small/Cash split.
# Add "HYBRID" here if you also want it for Balanced Advantage / Arbitrage
# (note: their debt / derivative portion would then fall into Cash & Others).
SPLIT_FUND_TYPES = {"EQUITY"}

# --- Reference file ------------------------------------------------------------
REF_SHEET = 0
REF_ISIN_HEADER_KEYWORD = "isin"              # header text that identifies the ISIN column
REF_CATEGORY_HEADER_KEYWORD = "categorization"  # header text that identifies the category column
# category text in reference file (lower-case) -> label used in output
CATEGORY_MAP = {
    "large cap": "Large Cap",
    "mid cap":   "Mid Cap",
    "small cap": "Small Cap",
}
CATEGORY_ORDER = ["Large Cap", "Mid Cap", "Small Cap"]

# --- Portfolio file layout (detected automatically from the header row) -------
SHEET_HEADER_SEARCH_ROWS = 40                 # header row must be within the first N rows
NAME_HEADER_KEYWORDS = ["nameoftheinstrument", "nameofinstrument", "instrument", "securityname", "issuer"]
ISIN_HEADER_KEYWORD = "isin"
PCT_HEADER_KEYWORD = "%to"                    # matches "% to Net Assets", "% to NAV", ...

# Rows whose name STARTS WITH any of these (spaces/dashes/case ignored) are not holdings
EXCLUDE_NAME_STARTS = ["subtotal", "total", "grandtotal", "netreceivable", "netcurrentasset",
                       "netpayable", "netassets"]
# TREPS / repo lines are real holdings for debt & liquid funds (an Overnight fund is 100%
# TREPS), but not for equity funds. Choose per fund type whether they count in the top 10.
CASH_LIKE_STARTS = ["treps", "triparty", "reverserepo", "cblo"]
CASH_LIKE_IN_TOP10_BY_TYPE = {"EQUITY": False, "HYBRID": False, "DEBT": True, "UNKNOWN": False}

# =============================================================================
# 2. HELPERS
# =============================================================================

WARNINGS = []


def warn(msg):
    if msg not in WARNINGS:
        WARNINGS.append(msg)


def norm(text):
    """lower-case and strip every whitespace / dash type."""
    if text is None or (not isinstance(text, str) and pd.isna(text)):
        return ""
    return re.sub(r"[\s\-\u2010-\u2015\u2212]+", "", str(text).lower())


def is_text(v):
    return isinstance(v, str) and v.strip() != ""


def clean_isin(v):
    if not is_text(v):
        return None
    s = v.strip().upper()
    return s if re.fullmatch(r"[A-Z]{2}[A-Z0-9]{9}[0-9]", s) else None


def find_dir(path):
    """Return `path`, or a same-named folder with different capitalisation."""
    if os.path.isdir(path):
        return path
    parent, name = os.path.split(path)
    if os.path.isdir(parent):
        for entry in os.listdir(parent):
            if entry.lower() == name.lower() and os.path.isdir(os.path.join(parent, entry)):
                return os.path.join(parent, entry)
    return path


# =============================================================================
# 3. REFERENCE FILE
# =============================================================================

def load_reference():
    if not os.path.exists(REFERENCE_FILE):
        raise FileNotFoundError(f"Reference file not found -> {REFERENCE_FILE}")
    raw = pd.read_excel(REFERENCE_FILE, sheet_name=REF_SHEET, header=None, dtype=object)
    hdr = isin_col = cat_col = None
    for r in range(min(15, len(raw))):
        cells = [norm(v) for v in raw.iloc[r]]
        i = next((c for c, t in enumerate(cells) if REF_ISIN_HEADER_KEYWORD in t), None)
        k = next((c for c, t in enumerate(cells) if REF_CATEGORY_HEADER_KEYWORD in t), None)
        if i is not None and k is not None:
            hdr, isin_col, cat_col = r, i, k
            break
    if hdr is None:
        raise ValueError("Could not find the ISIN / Categorization header row in the reference file")

    mapping = {}
    unknown = set()
    for isin, cat in zip(raw.iloc[hdr + 1:, isin_col], raw.iloc[hdr + 1:, cat_col]):
        code = clean_isin(isin)
        if not code or not is_text(cat):
            continue
        label = CATEGORY_MAP.get(cat.strip().lower())
        if label is None:
            unknown.add(cat.strip())
            continue
        mapping[code] = label
    if unknown:
        warn(f"Reference file has category values not in CATEGORY_MAP (ignored): {sorted(unknown)}")
    counts = pd.Series(mapping).value_counts().to_dict()
    print(f"Reference loaded: {len(mapping)} ISINs {counts}")
    return mapping


# =============================================================================
# 4. PORTFOLIO FILES
# =============================================================================

def parse_portfolio_sheet(raw):
    """raw = DataFrame (header=None). Returns dict or None if no holdings table is found."""
    nrows = len(raw)

    # ---- fund title & statement date ----
    title = subtitle = None
    for r in range(min(10, nrows)):
        for v in raw.iloc[r]:
            if not is_text(v) or v.strip() == "-":
                continue
            if "portfoliostatement" in norm(v):
                subtitle = subtitle or v.strip()
            elif title is None:
                title = v.strip()
    # ---- header row ----
    hdr = name_col = pct_col = isin_col = None
    for r in range(min(SHEET_HEADER_SEARCH_ROWS, nrows)):
        cells = [norm(v) if is_text(v) else "" for v in raw.iloc[r]]
        n = next((c for c, t in enumerate(cells) if t and any(k in t for k in NAME_HEADER_KEYWORDS)), None)
        p = next((c for c, t in enumerate(cells) if PCT_HEADER_KEYWORD in t), None)
        if n is not None and p is not None and n != p:
            hdr, name_col, pct_col = r, n, p
            isin_col = next((c for c, t in enumerate(cells) if ISIN_HEADER_KEYWORD in t), None)
            break
    if hdr is None:
        return None

    # ---- holdings ----
    rows = []
    grand_total = None
    for r in range(hdr + 1, nrows):
        name = raw.iloc[r, name_col]
        if not is_text(name):
            continue
        key = norm(name)
        pct = pd.to_numeric(raw.iloc[r, pct_col], errors="coerce")
        if key.startswith("grandtotal"):
            grand_total = None if pd.isna(pct) else float(pct)
            break
        if pd.isna(pct):
            continue
        if any(key.startswith(s) for s in EXCLUDE_NAME_STARTS):
            continue
        if pct <= 0:
            continue
        isin = clean_isin(raw.iloc[r, isin_col]) if isin_col is not None else None
        rows.append({"name": name.strip(), "isin": isin, "pct": float(pct),
                     "cash_like": any(key.startswith(s) for s in CASH_LIKE_STARTS)})

    # values stored as fractions (Grand Total = 1) instead of percentages (= 100)
    scale = 100.0 if grand_total is not None and abs(grand_total - 1) < 0.05 else 1.0
    for h in rows:
        h["pct"] *= scale

    # merge duplicate lines of the same security
    merged = {}
    for h in rows:
        k = h["isin"] or norm(h["name"])
        if k in merged:
            merged[k]["pct"] += h["pct"]
        else:
            merged[k] = dict(h)
    holdings = list(merged.values())
    return {"title": title, "subtitle": subtitle, "holdings": holdings}


def identify_fund(title, sheet_name, file_name):
    """Return (rule index, display name, fund type) or (None, None, None)."""
    for source in (title, sheet_name, file_name):
        text = norm(source)
        if not text:
            continue
        for i, (disp, ftype, keys) in enumerate(FUND_RULES):
            if any(norm(k) in text for k in keys):
                return i, disp, ftype
    return None, None, None


def load_portfolios():
    folder = find_dir(PORTFOLIO_DIR)
    if not os.path.isdir(folder):
        raise FileNotFoundError(f"Portfolio folder not found -> {PORTFOLIO_DIR}")
    files = sorted(f for f in os.listdir(folder)
                   if f.lower().endswith((".xlsx", ".xlsm")) and not f.startswith("~$"))
    if not files:
        raise FileNotFoundError(f"No Excel files found in {folder}")

    funds = []
    for f in files:
        path = os.path.join(folder, f)
        try:
            sheets = pd.read_excel(path, sheet_name=None, header=None, dtype=object)
        except Exception as e:                       # corrupt / locked file
            warn(f"{f}: could not be read ({e})")
            continue
        for sname, raw in sheets.items():
            parsed = parse_portfolio_sheet(raw)
            if parsed is None or not parsed["holdings"]:
                warn(f"{f} [{sname}]: no holdings table found - skipped")
                continue
            idx, disp, ftype = identify_fund(parsed["title"], sname, os.path.splitext(f)[0])
            if idx is None:
                warn(f"{f} [{sname}]: fund '{parsed['title']}' did not match any FUND_RULES row - "
                     f"top 10 is produced but no market-cap split")
                idx, disp, ftype = 999, parsed["title"] or sname, "UNKNOWN"
            funds.append({"file": f, "sheet": sname, "rule_idx": idx, "rule_name": disp, "type": ftype,
                          "title": parsed["title"] or disp, "subtitle": parsed["subtitle"],
                          "holdings": parsed["holdings"]})
            print(f"  {f}  [{sname}]  ->  {parsed['title']}  =>  {disp} ({ftype}), "
                  f"{len(parsed['holdings'])} holdings")
    # duplicates (two files for the same fund)
    seen = {}
    for fd in funds:
        if fd["type"] != "UNKNOWN":
            if fd["rule_name"] in seen:
                warn(f"Two portfolio files matched '{fd['rule_name']}': {seen[fd['rule_name']]} and {fd['file']}")
            seen[fd["rule_name"]] = fd["file"]
    funds.sort(key=lambda d: (d["rule_idx"], d["file"]))
    return funds


# =============================================================================
# 5. CALCULATIONS
# =============================================================================

def top_holdings(fd, n=TOP_N):
    keep_cash = CASH_LIKE_IN_TOP10_BY_TYPE.get(fd["type"], False)
    pool = [h for h in fd["holdings"] if keep_cash or not h["cash_like"]]
    return sorted(pool, key=lambda h: -h["pct"])[:n]


def market_cap_split(holdings, ref):
    sums = {c: 0.0 for c in CATEGORY_ORDER}
    unmatched = []
    for h in holdings:
        if not h["isin"]:
            continue
        cat = ref.get(h["isin"])
        if cat:
            sums[cat] += h["pct"]
        elif h["isin"].startswith("INE"):            # Indian-listed security not in reference
            unmatched.append(h)
    return sums, unmatched


# =============================================================================
# 6. OUTPUT
# =============================================================================

FONT = "Arial"
HDR_FILL = PatternFill("solid", start_color="1F3864")
TOT_FILL = PatternFill("solid", start_color="D9E1F2")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
PCT_FMT = "0.00"


def style(c, bold=False, color="000000", fill=None, align=None, fmt=None, wrap=False):
    c.font = Font(name=FONT, size=10, bold=bold, color=color)
    if fill:
        c.fill = fill
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    c.border = BORDER
    if fmt:
        c.number_format = fmt


def short_title(title):
    t = title
    for p in STRIP_FROM_SHEET_NAMES:
        t = t.replace(p, "")
    return t.strip()


def unique_sheet_name(title, used):
    base = re.sub(r"[\[\]:*?/\\]", "", short_title(title))[:31] or "Fund"
    name, i = base, 2
    while name.lower() in used:
        suffix = f" ({i})"
        name = base[:31 - len(suffix)] + suffix
        i += 1
    used.add(name.lower())
    return name


def heading(ws, fd, line2):
    ws["A1"] = fd["title"]
    ws["A1"].font = Font(name=FONT, size=12, bold=True)
    ws["A2"] = line2
    ws["A2"].font = Font(name=FONT, size=9, italic=True)


def header_row(ws, row, labels):
    for c, text in enumerate(labels, 1):
        ws.cell(row, c, text)
        style(ws.cell(row, c), True, "FFFFFF", HDR_FILL, "center")


def write_top10(funds):
    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for fd in funds:
        ws = wb.create_sheet(unique_sheet_name(fd["title"], used))
        heading(ws, fd, f"Top {TOP_N} holdings - {fd['subtitle'] or ''}".rstrip(" -"))
        header_row(ws, 4, ["Rank", "Name of the Instrument", "% to Net Assets"])
        top = top_holdings(fd)
        for i, h in enumerate(top, 1):
            r = 4 + i
            ws.cell(r, 1, i)
            ws.cell(r, 2, h["name"])
            ws.cell(r, 3, h["pct"])
            style(ws.cell(r, 1), align="center")
            style(ws.cell(r, 2))
            style(ws.cell(r, 3), align="right", fmt=PCT_FMT)
        last = 4 + len(top)
        tr = last + 1
        ws.cell(tr, 2, f"Total of top {len(top)}")
        ws.cell(tr, 3, f"=SUM(C5:C{last})")
        style(ws.cell(tr, 1), fill=TOT_FILL)
        style(ws.cell(tr, 2), True, fill=TOT_FILL)
        style(ws.cell(tr, 3), True, fill=TOT_FILL, align="right", fmt=PCT_FMT)
        ws.column_dimensions["A"].width = 8
        ws.column_dimensions["B"].width = 60
        ws.column_dimensions["C"].width = 18
    wb.save(TOP10_OUTPUT)


def write_split(funds, ref):
    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    n_sheets = 0
    for fd in funds:
        if fd["type"] not in SPLIT_FUND_TYPES:
            continue
        sums, unmatched = market_cap_split(fd["holdings"], ref)
        ws = wb.create_sheet(unique_sheet_name(fd["title"], used))
        n_sheets += 1
        heading(ws, fd, f"Market capitalisation split - {fd['subtitle'] or ''}".rstrip(" -"))
        header_row(ws, 4, ["Category", "% to Net Assets"])
        for i, cat in enumerate(CATEGORY_ORDER, 1):
            ws.cell(4 + i, 1, cat)
            ws.cell(4 + i, 2, sums[cat])
            style(ws.cell(4 + i, 1))
            style(ws.cell(4 + i, 2), align="right", fmt=PCT_FMT)
        first, last = 5, 4 + len(CATEGORY_ORDER)
        cr = last + 1
        ws.cell(cr, 1, "Cash & Others")
        ws.cell(cr, 2, f"={TOTAL_ASSETS_PCT}-SUM(B{first}:B{last})")
        style(ws.cell(cr, 1))
        style(ws.cell(cr, 2), align="right", fmt=PCT_FMT)
        tr = cr + 1
        ws.cell(tr, 1, "Total")
        ws.cell(tr, 2, f"=SUM(B{first}:B{cr})")
        style(ws.cell(tr, 1), True, fill=TOT_FILL)
        style(ws.cell(tr, 2), True, fill=TOT_FILL, align="right", fmt=PCT_FMT)

        cash_val = TOTAL_ASSETS_PCT - sum(sums.values())
        if cash_val < -0.01:
            warn(f"{fd['title']}: Large+Mid+Small = {sum(sums.values()):.2f}% which is above {TOTAL_ASSETS_PCT}% "
                 f"- Cash & Others is negative")

        if unmatched:
            r0 = tr + 3
            ws.cell(r0, 1, "Unclassified equity holdings (ISIN not in reference file - counted in Cash & Others)")
            ws.cell(r0, 1).font = Font(name=FONT, size=10, bold=True)
            header_row(ws, r0 + 1, ["Name of the Instrument", "% to Net Assets", "ISIN"])
            for j, h in enumerate(sorted(unmatched, key=lambda x: -x["pct"]), 1):
                r = r0 + 1 + j
                ws.cell(r, 1, h["name"]); ws.cell(r, 2, h["pct"]); ws.cell(r, 3, h["isin"])
                style(ws.cell(r, 1)); style(ws.cell(r, 2), align="right", fmt=PCT_FMT); style(ws.cell(r, 3))
            tot_un = sum(h["pct"] for h in unmatched)
            warn(f"{fd['title']}: {len(unmatched)} equity holding(s) not found in reference file "
                 f"({tot_un:.2f}% of assets, shown in Cash & Others): "
                 + ", ".join(h["name"] for h in unmatched))
        ws.column_dimensions["A"].width = 44
        ws.column_dimensions["B"].width = 18
        ws.column_dimensions["C"].width = 16
    if n_sheets == 0:
        ws = wb.create_sheet("No equity funds")
        ws["A1"] = "No equity funds were found in the portfolio folder."
    wb.save(SPLIT_OUTPUT)
    return n_sheets


# =============================================================================
# 7. MAIN
# =============================================================================

def main():
    ref = load_reference()
    print(f"Reading portfolios from {find_dir(PORTFOLIO_DIR)}")
    funds = load_portfolios()
    if not funds:
        raise ValueError("No usable portfolio found")

    write_top10(funds)
    n = write_split(funds, ref)
    print(f"\nSaved {len(funds)} fund sheet(s) -> {TOP10_OUTPUT}")
    print(f"Saved {n} equity fund sheet(s) -> {SPLIT_OUTPUT}")
    if WARNINGS:
        print(f"\n{len(WARNINGS)} WARNING(S):")
        for w in WARNINGS:
            print("  - " + w)
    else:
        print("No warnings.")


if __name__ == "__main__":
    try:
        main()
    except (FileNotFoundError, ValueError, PermissionError) as err:
        sys.exit(f"ERROR: {err}")
